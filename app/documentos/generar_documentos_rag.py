from pathlib import Path

from docx import Document
from docx.shared import Pt


CARPETA_DOCUMENTOS = Path(__file__).resolve().parent


    """
    Creación de documentos necesarios para el proceso RAG, lo hacemos desde un documento .py para asegurar la correcta reproducción del proyecto.
    """

DOCUMENTOS = {
    "A_aclaraciones_bancarias.docx": {
        "titulo": "Atención de aclaraciones bancarias",
        "codigo": "A",
        "contenido": [
            (
                "1. Identificación del proceso",
                """
Código de proceso: A
Nombre del proceso: Atención de aclaraciones bancarias
Institución: BANCOMEX
Tipo de documento: Procedimiento operativo interno

El proceso de atención de aclaraciones bancarias forma parte de los procedimientos internos de BANCOMEX para revisar inconformidades relacionadas con operaciones registradas en productos financieros del cliente. Este procedimiento permite registrar, analizar y resolver solicitudes donde exista duda sobre cargos, abonos, comisiones, pagos aplicados, transferencias, retiros, depósitos u otros movimientos reflejados en los sistemas internos del banco.
                """,
            ),
            (
                "2. Objetivo",
                """
El objetivo del proceso es asegurar que toda aclaración bancaria sea atendida de forma ordenada, trazable y respaldada por evidencia operativa. La revisión debe realizarse con base en información disponible en sistemas autorizados de BANCOMEX, evitando conclusiones sin soporte documental. El proceso busca proteger la integridad de la información operativa, mantener seguimiento oportuno de cada caso y comunicar una respuesta clara sobre la procedencia o no procedencia de la aclaración.
                """,
            ),
            (
                "3. Alcance",
                """
Este procedimiento aplica cuando un cliente solicita revisión sobre una operación que considera incorrecta, desconocida, incompleta o no aplicada. También aplica cuando un área interna detecta una posible inconsistencia relacionada con cargos, abonos, comisiones o conciliaciones. No corresponde a este proceso atender cancelaciones de productos, actualización de datos del cliente, quejas internas sobre atención recibida ni fallas técnicas que requieran soporte operativo especializado. Cuando se identifique alguno de esos casos, la solicitud debe canalizarse al proceso correspondiente.
                """,
            ),
            (
                "4. Entradas requeridas",
                """
Para iniciar una aclaración se requiere contar con el identificador del cliente, producto relacionado, fecha aproximada de la operación, monto, canal donde se realizó el movimiento y descripción clara de la inconformidad. La información debe permitir ubicar la operación en los sistemas internos. Cuando los datos recibidos sean insuficientes, el caso debe quedar pendiente de complemento antes de iniciar el análisis operativo. Ninguna aclaración debe cerrarse sin contar con elementos mínimos para revisar el movimiento cuestionado.
                """,
            ),
            (
                "5. Flujo operativo",
                """
El proceso inicia con el registro de la solicitud en el canal autorizado. El analista responsable clasifica el tipo de aclaración, identifica el producto asociado y revisa si existe una solicitud duplicada sobre la misma operación. Posteriormente consulta los sistemas internos para validar si el movimiento existe, si fue aplicado correctamente, si existen reversos, comisiones asociadas, operaciones duplicadas o registros pendientes de conciliación.

Si la evidencia confirma que la operación fue aplicada correctamente, el caso se documenta como no procedente y se comunica la explicación correspondiente. Si se detecta una inconsistencia operativa, el analista registra la causa, solicita la corrección al área que corresponda y da seguimiento hasta confirmar el ajuste o la resolución del movimiento.
                """,
            ),
            (
                "6. Validaciones obligatorias",
                """
Antes de emitir una respuesta, debe validarse la identidad del cliente conforme al canal utilizado, confirmar que el producto pertenece al cliente, revisar que la solicitud esté completa y verificar que no exista una aclaración previa sobre el mismo movimiento. También debe validarse si la operación se encuentra dentro del periodo permitido para revisión interna. Si el caso está fuera del periodo regular, debe solicitarse autorización o validación adicional antes de continuar.
                """,
            ),
            (
                "7. Escalamiento y cierre",
                """
Una aclaración debe escalarse cuando involucra posible fraude, montos altos, inconsistencias entre sistemas, afectación a varios clientes o falta de evidencia suficiente para resolver en primer nivel. El escalamiento debe incluir motivo, evidencia revisada, sistemas consultados y área destino. No debe prometerse una resolución antes de recibir confirmación del área responsable de atender el escalamiento.

El cierre puede clasificarse como procedente, no procedente, pendiente de información o escalado. En todos los casos debe quedar evidencia del análisis realizado, estado final, motivo de la decisión y acción correctiva cuando aplique. La respuesta final debe ser clara, verificable y limitada a la información confirmada durante la revisión del caso.
                """,
            ),
        ],
    },
    "B_cancelacion_productos.docx": {
        "titulo": "Cancelación de productos financieros",
        "codigo": "B",
        "contenido": [
            (
                "1. Identificación del proceso",
                """
Código de proceso: B
Nombre del proceso: Cancelación de productos financieros
Institución: BANCOMEX
Tipo de documento: Procedimiento operativo interno

El proceso de cancelación de productos financieros establece los lineamientos internos para atender solicitudes de cierre operativo de productos activos del cliente. Puede aplicar sobre cuentas, tarjetas, servicios asociados, productos transaccionales u otros instrumentos financieros administrados por BANCOMEX, siempre que el producto sea elegible para cierre y se cumplan las validaciones previas establecidas.
                """,
            ),
            (
                "2. Objetivo",
                """
El objetivo es asegurar que la cancelación de productos financieros se realice de manera controlada, documentada y sin afectar obligaciones pendientes del cliente o del banco. La cancelación debe ejecutarse únicamente cuando existan condiciones operativas válidas, autorización suficiente y evidencia de que no hay bloqueos, cargos en proceso, investigaciones activas, operaciones retenidas o restricciones internas que impidan completar el cierre.
                """,
            ),
            (
                "3. Alcance",
                """
Este procedimiento aplica cuando el cliente solicita cancelar un producto financiero o cuando un área interna requiere cerrar un producto por una condición operativa documentada. No cubre aclaraciones por cargos o movimientos, actualización de datos personales, quejas internas ni incidencias técnicas. Si durante la revisión se identifica una inconformidad sobre movimientos, comisiones o cargos, el caso debe canalizarse al proceso de atención de aclaraciones bancarias.
                """,
            ),
            (
                "4. Entradas requeridas",
                """
Para iniciar la cancelación se requiere identificar al cliente, producto a cancelar, canal de solicitud, motivo de cancelación y validación de identidad o autorización interna. Cuando aplique, deben revisarse documentos soporte, consentimiento del cliente, saldos vigentes, productos relacionados y condiciones contractuales. Si falta documentación o existen restricciones operativas, el caso debe permanecer pendiente hasta completar la información requerida.
                """,
            ),
            (
                "5. Flujo operativo",
                """
El proceso inicia con el registro de la solicitud en el canal autorizado. El responsable valida que el producto exista, que pertenezca al cliente y que la solicitud esté autorizada. Después revisa si el producto tiene saldo pendiente, cargos en tránsito, operaciones retenidas, bloqueos preventivos, reclamaciones abiertas o dependencias con otros productos.

Si no existen restricciones, se procede con el cierre operativo y se documenta la fecha, canal, producto cancelado y evidencia de validación. Si se detecta una condición que impide cancelar, se informa el motivo y se registra la acción necesaria para continuar. Ningún producto debe cerrarse cuando exista una operación pendiente que pueda generar afectaciones posteriores.
                """,
            ),
            (
                "6. Validaciones obligatorias",
                """
Antes de cancelar un producto, debe validarse la identidad del cliente, el estado del producto, saldos, movimientos pendientes, bloqueos, cargos en proceso y obligaciones asociadas. También debe revisarse si existen productos vinculados, domiciliaciones, servicios recurrentes o restricciones regulatorias. La validación debe quedar documentada en el expediente operativo del caso.
                """,
            ),
            (
                "7. Escalamiento y cierre",
                """
El caso debe escalarse cuando exista saldo pendiente no conciliado, bloqueo preventivo, investigación activa, documentación incompleta, diferencia entre sistemas o restricción contractual. El escalamiento debe incluir el motivo, la evidencia consultada y la condición que impide completar el cierre.

El cierre puede ser exitoso, rechazado por condición operativa, pendiente de información o escalado a una mesa resolutora. La respuesta debe indicar el estado del cierre, el motivo operativo y los pasos requeridos cuando la cancelación no pueda completarse. Toda cancelación debe conservar evidencia suficiente para auditoría interna.
                """,
            ),
        ],
    },
    "C_escalamiento_incidencias.docx": {
        "titulo": "Escalamiento de incidencias operativas",
        "codigo": "C",
        "contenido": [
            (
                "1. Identificación del proceso",
                """
Código de proceso: C
Nombre del proceso: Escalamiento de incidencias operativas
Institución: BANCOMEX
Tipo de documento: Procedimiento operativo interno

El proceso de escalamiento de incidencias operativas define la forma en que BANCOMEX registra, clasifica y canaliza fallas, bloqueos o comportamientos no esperados que impiden completar una operación interna. Este procedimiento busca mantener continuidad operativa, reducir tiempos de afectación y asegurar que cada incidencia tenga evidencia, prioridad y responsable asignado.
                """,
            ),
            (
                "2. Objetivo",
                """
El objetivo es atender incidencias que no pueden resolverse en primer nivel o que requieren intervención de un equipo especializado. El proceso permite identificar el sistema afectado, documentar el impacto, clasificar la prioridad y enviar el caso al área con capacidad de resolverlo. La atención debe basarse en evidencia clara y no en suposiciones sobre la causa del error.
                """,
            ),
            (
                "3. Alcance",
                """
Este procedimiento aplica cuando un usuario interno reporta errores en sistemas, bloqueos de operación, inconsistencias entre plataformas, rechazos no esperados, lentitud crítica o indisponibilidad funcional que afecta procesos bancarios. No sustituye los procesos de aclaraciones, cancelación de productos, actualización de datos o quejas internas, aunque puede relacionarse con ellos cuando una falla impide completar dichos procesos.
                """,
            ),
            (
                "4. Entradas requeridas",
                """
Para registrar una incidencia se requiere descripción del problema, sistema afectado, usuario reportante, fecha y hora aproximada, operación que se intentaba realizar, mensaje de error cuando exista y evidencia disponible. También debe indicarse si la afectación corresponde a un solo usuario, una sucursal, un proceso específico o varias áreas operativas. Sin esta información, la clasificación puede quedar incompleta y retrasar el análisis.
                """,
            ),
            (
                "5. Flujo operativo",
                """
El proceso inicia cuando la mesa de ayuda operativa recibe el reporte. El primer nivel revisa si existe una solución conocida, si el error es reproducible y si el usuario cuenta con permisos adecuados. Si la incidencia se resuelve con una acción documentada, se registra la solución aplicada y se solicita confirmación del usuario.

Cuando la incidencia no puede resolverse en primer nivel, se clasifica por impacto y urgencia. Después se escala al equipo correspondiente con evidencia, pasos realizados, sistema afectado y prioridad sugerida. El área receptora debe analizar la causa, aplicar corrección o definir una ruta alterna autorizada para continuar la operación.
                """,
            ),
            (
                "6. Validaciones obligatorias",
                """
Antes de escalar, debe confirmarse que el reporte no corresponde a un error de captura, falta de permisos, información incompleta o uso incorrecto del sistema. También debe validarse si existen reportes similares abiertos. Cuando varios usuarios reportan la misma afectación, la incidencia debe tratarse como evento de mayor prioridad y documentarse el alcance operativo.
                """,
            ),
            (
                "7. Escalamiento y cierre",
                """
La incidencia debe escalarse cuando no existe solución conocida, afecta continuidad operativa, impacta atención a clientes, compromete varios procesos o presenta diferencias entre sistemas. El escalamiento debe incluir evidencia, pasos realizados, alcance de afectación y prioridad sugerida.

El cierre requiere documentar causa identificada, acción aplicada, área que resolvió y confirmación del usuario interno. No debe cerrarse una incidencia sin evidencia de recuperación o sin una ruta alterna autorizada. Si la causa no fue identificada, el caso debe mantenerse en seguimiento o documentarse como pendiente de análisis especializado.
                """,
            ),
        ],
    },
    "D_actualizacion_datos_cliente.docx": {
        "titulo": "Actualización de datos del cliente",
        "codigo": "D",
        "contenido": [
            (
                "1. Identificación del proceso",
                """
Código de proceso: D
Nombre del proceso: Actualización de datos del cliente
Institución: BANCOMEX
Tipo de documento: Procedimiento operativo interno

El proceso de actualización de datos del cliente regula la modificación de información registrada en los sistemas internos de BANCOMEX. Incluye cambios en datos generales, fiscales, de contacto, identificación y otros atributos necesarios para mantener expedientes actualizados y consistentes entre plataformas operativas.
                """,
            ),
            (
                "2. Objetivo",
                """
El objetivo es asegurar que cualquier modificación de datos del cliente se realice con documentación válida, autorización suficiente, trazabilidad y apego a reglas internas de integridad de información. El proceso busca evitar cambios no autorizados, duplicidad de registros, inconsistencias entre sistemas y afectaciones a productos o servicios vinculados al cliente.
                """,
            ),
            (
                "3. Alcance",
                """
Este procedimiento aplica para cambios de domicilio, teléfono, correo electrónico, actividad económica, datos fiscales, identificaciones, nombre registrado, información de contacto o corrección de datos generales. No cubre cancelación de productos, aclaraciones sobre movimientos, quejas internas ni incidencias técnicas. Sin embargo, puede relacionarse con otros procesos cuando un dato incorrecto impide completar una operación.
                """,
            ),
            (
                "4. Entradas requeridas",
                """
Para iniciar la actualización se requiere identificar al cliente, dato a modificar, valor actual, nuevo valor solicitado, canal de origen y documentación soporte. En cambios sensibles, como nombre, identificación, datos fiscales o actividad económica, se requiere revisión adicional conforme a reglas internas. Si la documentación es insuficiente, la solicitud queda pendiente hasta recibir complemento.
                """,
            ),
            (
                "5. Flujo operativo",
                """
El proceso inicia con el registro de la solicitud en el portal interno correspondiente. El analista revisa el tipo de dato a modificar y valida si el cambio puede realizarse con la documentación presentada. Posteriormente verifica consistencia entre los datos actuales y los nuevos, así como posibles impactos en productos relacionados.

Cuando la información es correcta, se actualiza el dato en el sistema maestro correspondiente y se valida si debe replicarse a otros sistemas. Si existe inconsistencia documental, posible duplicidad de cliente o conflicto entre plataformas, el caso debe escalarse antes de aplicar cualquier modificación.
                """,
            ),
            (
                "6. Validaciones obligatorias",
                """
Antes de actualizar datos, debe validarse identidad del cliente, vigencia de documentos, coincidencia entre documentos y datos registrados, reglas internas para cambios sensibles y existencia de registros duplicados. También debe confirmarse que el cambio no contradiga restricciones activas o información validada previamente por otra área de control.
                """,
            ),
            (
                "7. Escalamiento y cierre",
                """
El caso debe escalarse cuando exista inconsistencia documental, posible suplantación, duplicidad de cliente, conflicto entre sistemas maestros o cambio sensible sin respaldo suficiente. El escalamiento debe incluir el dato solicitado, la inconsistencia detectada y la evidencia revisada.

El cierre ocurre cuando el dato queda actualizado, rechazado por falta de soporte o escalado para validación adicional. La respuesta debe indicar el estado del cambio, la razón operativa y si existe alguna acción pendiente. No debe modificarse información del cliente sin evidencia suficiente o autorización válida.
                """,
            ),
        ],
    },
    "E_quejas_internas.docx": {
        "titulo": "Gestión de quejas internas",
        "codigo": "E",
        "contenido": [
            (
                "1. Identificación del proceso",
                """
Código de proceso: E
Nombre del proceso: Gestión de quejas internas
Institución: BANCOMEX
Tipo de documento: Procedimiento operativo interno

El proceso de gestión de quejas internas establece la forma en que BANCOMEX registra, analiza y da seguimiento a inconformidades relacionadas con atención, operación, cumplimiento, tiempos de respuesta o comportamiento interno. Este procedimiento permite documentar los casos, asignar responsables y dar seguimiento hasta contar con una conclusión formal.
                """,
            ),
            (
                "2. Objetivo",
                """
El objetivo es asegurar que toda queja interna sea tratada con trazabilidad, imparcialidad y evidencia suficiente. El proceso busca identificar causas operativas, detectar reincidencias, proponer acciones correctivas cuando correspondan y mantener registro de las decisiones tomadas. La atención de una queja no debe basarse en opiniones no verificadas ni en información ajena al expediente del caso.
                """,
            ),
            (
                "3. Alcance",
                """
Este procedimiento aplica cuando un colaborador, área interna o canal autorizado reporta una inconformidad sobre atención recibida, incumplimiento de procedimientos, demoras injustificadas, trato inadecuado, falta de seguimiento o posibles desviaciones operativas. No reemplaza los procesos de aclaraciones, cancelación de productos, actualización de datos o incidencias técnicas, aunque puede recibir casos derivados cuando exista una inconformidad formal sobre la gestión.
                """,
            ),
            (
                "4. Entradas requeridas",
                """
Para registrar una queja interna se requiere descripción del caso, área involucrada, fecha, canal de origen, evidencia disponible, impacto observado y datos del responsable que reporta. Cuando la información es insuficiente, se solicita complemento antes de iniciar análisis. Si la queja se relaciona con un proceso operativo específico, debe incluirse el folio o referencia del caso asociado.
                """,
            ),
            (
                "5. Flujo operativo",
                """
La queja se registra en el sistema interno correspondiente. La instancia encargada revisa la información inicial, valida que no exista duplicidad y asigna un responsable de análisis. Según la naturaleza del caso, puede requerirse revisión de expedientes, tiempos de atención, comunicaciones internas, cumplimiento de procedimiento o participación de otra área.

Durante el análisis se documentan hallazgos, evidencias revisadas y acciones solicitadas. Si se requiere respuesta de otra área, se establece un plazo de atención y se da seguimiento hasta recibir información suficiente. La queja no debe cerrarse sin conclusión documentada.
                """,
            ),
            (
                "6. Validaciones obligatorias",
                """
Antes de cerrar una queja, debe validarse que el caso esté completo, que la evidencia haya sido revisada, que el área involucrada haya tenido oportunidad de responder y que la conclusión esté respaldada por información verificable. También debe revisarse si existen quejas previas similares, reincidencia o impacto en clientes o procesos críticos.
                """,
            ),
            (
                "7. Escalamiento y cierre",
                """
La queja debe escalarse cuando exista posible incumplimiento normativo, afectación a clientes, reincidencia, conflicto entre áreas, falta de respuesta del responsable asignado o posible desviación de controles internos. El escalamiento debe quedar documentado con motivo, evidencia y área destino.

El cierre debe incluir conclusión, acciones correctivas si aplican, responsable de seguimiento y fecha de resolución. No deben afirmarse sanciones, responsables o decisiones sin soporte documental. Si la información no permite concluir el caso, debe permanecer en seguimiento hasta contar con evidencia suficiente.
                """,
            ),
        ],
    },
}



    """
    Definimos las caracteristicas que serán uasadas para crear los documentos Word.
    """
def limpiar_texto(texto: str) -> str:
    return " ".join(texto.strip().split())


def configurar_estilo(documento: Document) -> None:
    estilo_normal = documento.styles["Normal"]
    estilo_normal.font.name = "Arial"
    estilo_normal.font.size = Pt(10.5)


def crear_documento(nombre_archivo: str, datos: dict) -> None:
    documento = Document()
    configurar_estilo(documento)

    documento.add_heading(datos["titulo"], level=1)

    documento.add_paragraph(f"Código de proceso: {datos['codigo']}")
    documento.add_paragraph(f"Nombre del proceso: {datos['titulo']}")
    documento.add_paragraph("Institución: BANCOMEX")
    documento.add_paragraph("Tipo de documento: Procedimiento operativo interno")

    for titulo, texto in datos["contenido"]:
        documento.add_heading(titulo, level=2)
        documento.add_paragraph(limpiar_texto(texto))

    ruta_salida = CARPETA_DOCUMENTOS / nombre_archivo
    documento.save(ruta_salida)
    print(f"Documento generado: {ruta_salida}")


def main() -> None:
    for nombre_archivo, datos in DOCUMENTOS.items():
        crear_documento(nombre_archivo, datos)


if __name__ == "__main__":
    main()